#!/usr/bin/env python3
r"""
zzbuild_testbook_status.py — ONE-TIME seed of the IDEAL testbook_status.xlsx spec (2026-06-16).

testbook_status.xlsx is the "god standard": the spec the testbook harness code must catch up to.
It is HAND-MAINTAINED after this (the human adds Approve/HumanFeedback; do NOT re-run this and
clobber that). Kept here only to document how the initial spec rows were seeded.
The 3 rows the human had already annotated are preserved verbatim (Approve/HumanFeedback/Status).
"""
import os
import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# self-locating (no machine path / username): scripts/_oneoff -> hk-config root -> data/.
_HKROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
P = os.path.join(_HKROOT, "data", "testbook_status.xlsx")

HDR = ["id", "Mod", "Feature", "Component", "Source", "GenExceptions", "DebugTokens",
       "TBElementType", "TBElementDetail", "Approve", "HumanFeedback", "Status"]

R = [
    # ---- versiontestMod ----
    ["VT-LOAD", "versiontest", "Load-confirm popup", "versiontest.1 fires every country, log player-gated", "script", "-",
     'debug_log=versiontest_log ("VTEST_BUILD <stamp>"); loc versiontest.1.d', "Log + BDD",
     "Log tab parses the VTEST_BUILD line -> auto-labels the report with the loaded build stamp; BDD: 'popup shown on load?' Y/N",
     "up", "fic error found uder log for this", "IN PROGRESS"],
    # ---- Top40EcoBoostMod ----
    ["ECO-CHAMP", "EcoBoost", "Champ seeding", "eco_pulse_spread champ N=tier/2 spread across states",
     "gen (gen_ecoboost_configs <- SELECT_ECOBOOST_FINAL_v2.csv)", "grain->primary staple farm; champ-excluded goods (wood/iron/oil/rubber/gold)",
     "$B$ header + eco_hit_champ + debug_log_scopes (state)", "Census ptab",
     "per-tag: champ good, target N, built count, state list; flag under/over-build vs cap", "", "", "OPEN"],
    ["ECO-SUPPORT", "EcoBoost", "Support seeding", "eco_pulse_spread support N=tier/4 each good", "gen", "-",
     "$B$ + eco_hit_support + scope", "Census ptab", "per-tag: each support good, N, built count + states", "", "", "OPEN"],
    ["ECO-FLAVOUR", "EcoBoost", "Flavour seeding", "eco_seed_flavour +1 in best-GDP constructible state", "gen",
     "unbuildable flavour no-ops (can_construct gate); SHW vineyard buildable post-Oromia",
     "$B$ + eco_hit_flavour + eco_have_flavour + scope", "Census ptab + Timeline",
     "per-tag flavour built vs planned (eco_flavour_built/unbuilt cross-ref). GAP: log lacks STATE -> blocks history move (T48/T52)",
     "down", "Need refinemoent. Ptab need year filter too.Add year in debug log framework in tis event if you can source it from the logs.", "IN PROGRESS"],
    ["ECO-GENFLAV", "EcoBoost", "Generic anti-death-spiral flavour", "eco_generic_flavour wood+5 / staple+5 / food+1 every target", "hand",
     "auto-detects existing staple; food only if teched", "$B$ + eco_hit_flavour", "Census ptab", "wood/grain/food seed counts per tag", "", "", "OPEN"],
    ["ECO-CHAMPMOD", "EcoBoost", "Champ throughput+wage modifier", "add_modifier eco_champ_tp_{plant|log|ranch|fish|mine|ind} 20yr",
     "gen dispatch + hand modifiers", "champ good->category map (eco_champ_modifier_map.csv)",
     "(no log token) modifier name", "Modifier/standards check",
     "each tag gets the RIGHT category modifier (cross-ref eco_champ_modifier_map.csv); throughput<=0.5; has icon+loc", "", "", "OPEN"],
    ["ECO-SOLMOD", "EcoBoost", "Global SoL/edu modifier", "eco_sol_edu_global whole game", "hand", "-",
     "(modifier)", "Modifier check", "every target has it; state_standard_of_living_add=4, education_access_add=0.25", "", "", "OPEN"],
    ["ECO-YEARLY", "EcoBoost", "Yearly pulse spread", "eco.5 on_yearly_pulse re-runs champ/support (covers annexed states)", "gen",
     "overflow when state hits cap; ECONOSLOT when no slot", "eco_hit_champ/support per build; 'ECONOSLOT'", "Timeline ptab (BY YEAR)",
     "per-YEAR build additions per tag; needs a YEAR dimension in the debug-log framework (user)", "", "", "OPEN"],
    ["ECO-OILRUBBER", "EcoBoost", "Oil/rubber tech-triggered seed", "eco_oilrubber events + on_actions on tech unlock", "gen/hand",
     "fires only when oil/rubber tech researched", "(verify token) ECOOR / $B$", "Timeline / Log",
     "fires when oil/rubber tech unlocked. UNTESTED (Row36: test before test-branch)", "", "", "OPEN"],
    ["ECO-TRACKB", "EcoBoost", "Track-B starters (wood/iron)", "zz_eco_starters.txt history create_building", "gen (gen_ecoboost_starters)",
     "top-80% wood / top-40% iron states by pop", "(history - no runtime token)", "State-buildings census",
     "post-load state buildings vs expected per tag; bad-tuple scrub needs error.log (Row34)", "", "", "OPEN"],
    ["ECO-EASTAFR", "EcoBoost", "East Africa override", "zz_eco_starters_eastafrica.txt history", "hand (generator retired 2026-06-16)",
     "only EGY/SHW own EA states", "(history)", "State-buildings census", "ERITREA->EGY & OROMIA->SHW cotton present at start", "", "", "OPEN"],
    ["ECO-NOTIFY", "EcoBoost", "Player notification", "eco.4 visible event (player only)", "gen", "-",
     "(visible event)", "BDD", "'eco-seed notification shown to player?' Y/N", "", "", "OPEN"],
    # ---- MyDiploPlayMod ----
    ["MDP-WARS", "MyDiploPlay", "Day-1 curated wars", "zz_mdp_wars.txt history diplomatic_plays", "hand (gen frozen, T50)",
     "bad/GP-dragging wars commented out; return_state goals", "(history - no runtime token)", "Diplo/Timeline check",
     "each curated war present at start; full-state return_state goals; low infamy", "", "", "OPEN"],
    ["MDP-TRUCES", "MyDiploPlay", "All-expander truces", "zz_mdp_truces.txt history", "gen (gen_blob_plan) + hand edits",
     "region-bucket pairs + MANUAL (OMA-NEJ, OMA-SHW, SHW-WTU); CHT/KAB 24mo hand", "(history)", "Diplo-state ptab",
     "truce matrix at start; 120mo expires 1846.1 = blob-chain stop, covers whole window", "", "Not yes Build", "OPEN"],
    ["MDP-BLOB", "MyDiploPlay", "Runtime blob chain", "mdp.1 -> mdp_try_next_war target pick", "hand",
     "army_size<80% of attacker; biggest GDP; land-adjacent; truce-excluded",
     "mdp_try_next_war TARGET CHOSEN/WAR DECLARED + scope; MDP_CLAIM/MDP_HOMELAND markers", "Timeline ptab",
     "per attacker (by month): target chosen, #claims/#homelands, war declared", "", "", "OPEN"],
    ["MDP-NEPTIB", "MyDiploPlay", "NEP->TIB opium event", "mdp.2 gated on China opium_ban_authority", "hand",
     "fires only when China has the modifier", "mdp.2 FIRED / WAR DECLARED vs TIB", "Timeline / Log",
     "NEP attacks TIB only under China opium-ban penalty", "", "", "OPEN"],
    ["MDP-ANNEX", "MyDiploPlay", "Subject annexation", "mdp.3 dp_annex_subject, peace, <1851", "hand",
     "weakest DIRECT subject only", "mdp.3 SUBJECT CHOSEN / ANNEXATION PLAY DECLARED", "Timeline",
     "DAI/BUR/SIA/PAN annex a subject 1846-51", "", "", "OPEN"],
    ["MDP-DAISIA", "MyDiploPlay", "DAI/SIA non-interference + 1837 revert", "zz_mdp_favours + zz_mdp_relations + mdp.4 one-shot", "hand",
     "one-shot revert at 1837 (months=12)", "mdp.4 reverting DAI<->SIA non-interference (1837)", "Timeline / Diplo",
     "start: SIA owes DAI + DAI +50 SIA; reverts to 0/none in 1837", "", "", "OPEN"],
    ["MDP-INFAMY", "MyDiploPlay", "Infamy-decay modifiers", "mdp_infamy_decay 120mo heavy / 36mo standard", "hand",
     "heavy list (SEA+PAN/SOK/SHW/SAF/BRZ) vs rest; NEP/ARG=36mo", "(modifier) + debug mdp_infamy_armed_row", "Modifier check",
     "each expand tag gets the RIGHT decay duration", "", "", "OPEN"],
    ["MDP-CLAIMS", "MyDiploPlay", "Claims + homeland", "zz_mdp_claims.txt history", "gen (gen_blob_plan, human-edited)",
     "BIC/PRU homeland trimmed; SWE claims removed (human)", "(history)", "State check",
     "target states have add_claim + add_homeland at start", "", "", "OPEN"],
    ["MDP-SEEDPOP", "MyDiploPlay", "Seed pops", "zz_mdp_seed_pops.txt history", "gen (gen_blob_plan)",
     "NO european cultures (scrubbed 2026-06-16; EUROPEAN_CULTURES exclusion)", "(history)", "Pop check",
     "primary-culture pops in target states; ZERO british/dutch/etc", "", "", "OPEN"],
    # ---- ExpFightMod ----
    ["EXF-BOOST", "ExpFight", "War-footing combat boost", "warboost modifier (offense+30/def+50/morale/move/cost)", "hand", "-",
     "(modifier) + expfight on_actions debug", "Modifier check", "expand tags get warboost; verify offense=30/def=50", "", "", "OPEN"],
    ["EXF-BUMP", "ExpFight", "War-footing IG/morale bump", "warbump modifier (IG approval/pol_str/morale) 60mo", "hand", "-",
     "(modifier)", "Modifier check", "applied to expand tags at start", "", "", "OPEN"],
    ["EXF-LONGWAR", "ExpFight", "Long-war 120mo set", "expfight_is_long_war_country (11 tags)", "hand",
     "NEP/ARG removed (run out of local targets); strong powers excluded", "(trigger list)", "Config check",
     "right tags get 120mo footing vs 60mo", "", "", "OPEN"],
    # ---- ExpMktAccessMod ----
    ["EXM-LAW", "ExpMkt", "Interventionism switch", "expmkt_start_setup activate_law", "hand", "guarded vs laissez-faire (never downgrade)",
     "debug expmkt_start_row", "Log / BDD", "expand tags -> interventionism at start (unless already liberal)", "", "", "OPEN"],
    ["EXM-CAPTC", "ExpMkt", "Capital trade center", "create_building trade_center in capital", "hand", "only if missing",
     "debug expmkt_start_row", "State check", "capital has a trade center (covers inland tags)", "", "", "OPEN"],
    ["EXM-COAST", "ExpMkt", "Coastal infra seed", "expmkt_seed_coastal port + TC + logging(lvl1)", "hand",
     "can_construct-gated; logging capped at lvl1; FIRST runs 1837.6.1", "debug expmkt_coastal_seed_row + scope", "Timeline / State",
     "best coastal state seeded; first fire 1837.6.1 (months=17)", "", "", "IN PROGRESS"],
    ["EXM-PULSE", "ExpMkt", "4-year coastal pulse", "expmkt.1 yearly x4", "hand", "hard-stop after year 4",
     "debug expmkt_yearly_pulse_row", "Timeline (by year)", "pulse fires 1837.6 then yearly; catches mid-war coastal annexes", "", "", "OPEN"],
    ["EXM-INDMOD", "ExpMkt", "Industrialist boost modifier", "expmkt_industrialist_boost 48mo", "hand", "-",
     "(modifier)", "Modifier check", "applied to expand tags at start", "", "", "OPEN"],
    # ---- InfraTaxMod ----
    ["INF-STATE", "InfraTax", "State infra/tax/bureau traits", "map_data state_regions + state_traits", "gen (port_statedata)",
     "only USER-touched states (fingerprinted)", "(history/static - no token)", "State-modifier check",
     "seeded states carry the infra/tax/bureau modifiers", "", "", "OPEN"],
    ["INF-SEED", "InfraTax", "Infra building seeding (PLANNED)", "one-shot per-state ports/TC/university/rail", "planned",
     "tech-gated; MAX 1 per state; fire once", "(TBD)", "State check", "NEW TODO - not yet built", "", "", "OPEN"],
    # ---- BetterPopPromoMod ----
    ["BPP-GATES", "BetterPopPromo", "Tiered promotion gates", "t1/t2 gates", "script", "-",
     "(verify bpp tokens)", "Config check", "gates applied per tier; bpp_targets.csv wired", "", "", "OPEN"],
    ["BPP-CYCLE", "BetterPopPromo", "Yearly check + toast + conversion", "bpp.2 yearly / bpp.3 toast / bpp.4 BIC conversion", "script", "-",
     "(verify bpp debug)", "Timeline / BDD", "bpp.2 fires yearly; PRU/PAN/SOK toast; BIC pop-conversion test", "", "", "OPEN"],
]

LEGEND = [
    ("id", "Stable key linking this spec row to a testbook HTML section/anchor and to feedback."),
    ("Mod", "Which mod the feature belongs to."),
    ("Feature", "User-facing feature of the mod."),
    ("Component", "The concrete sub-piece of the feature that testbook must verify."),
    ("Source", "hard-coded | script-generated (name the gen_* + seed/data csv) | hand."),
    ("GenExceptions", "Exceptions/edge-cases to the generator logic (what the gen deliberately skips/special-cases)."),
    ("DebugTokens", "The debug_log markers / strings / loc keys emitted in the mod files that testbook parses to capture this component."),
    ("TBElementType", "Testbook element used: ptab (parsed table) | Timeline | Log | BDD | State/Modifier check | Census."),
    ("TBElementDetail", "IDEAL description of WHAT the testbook element should show / assert for this component."),
    ("Approve", "Human thumb up/down on the testbook capture for this row (HUMAN fills)."),
    ("HumanFeedback", "Human remarks on the capture / what to refine (HUMAN fills)."),
    ("Status", "Testbook conformance to this spec: OPEN=gap | IN PROGRESS=partial | CLOSED=captured+approved | CANCELLED."),
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "status"
    ws.append(HDR)
    for r in R:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate([12, 14, 26, 42, 22, 40, 46, 18, 64, 9, 40, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    lg = wb.create_sheet("Legend")
    lg.append(["Column", "Full meaning"])
    for row in LEGEND:
        lg.append(row)
    lg.column_dimensions["A"].width = 16
    lg.column_dimensions["B"].width = 110

    wb.save(P)
    from collections import Counter
    print(f"wrote testbook_status.xlsx: {len(R)} spec rows + Legend sheet")
    print("by mod:", dict(Counter(r[1] for r in R)))
    print("by status:", dict(Counter(r[11] for r in R)))


if __name__ == "__main__":
    main()
