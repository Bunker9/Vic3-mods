#!/usr/bin/env python3
"""githk_build_todo_xlsx — render hk-config/data/TODO.csv into a formatted, status
colour-coded TODO.xlsx (open=red, in_progress=yellow, done=green).

Self-locating: lives in hk-config/tools/, reads ../data/TODO.csv, writes ../data/TODO.xlsx.
Both CSV and XLSX are gitignored (private tracker) — this script is the only committed part,
so the sheet is always regenerable from the CSV. Re-run after editing TODO.csv.

Usage:  python tools/githk_build_todo_xlsx.py
"""
from __future__ import annotations
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATA = Path(__file__).resolve().parent.parent / "data"
CSV_PATH = DATA / "TODO.csv"
XLSX_PATH = DATA / "TODO.xlsx"

# status -> (fill, font colour)
STATUS_STYLE = {
    "open":        ("FFC7CE", "9C0006"),  # red
    "in_progress": ("FFEB9C", "9C6500"),  # yellow
    "done":        ("C6EFCE", "006100"),  # green
}
PRIORITY_FONT = {"high": "C00000", "mid": "808000", "low": "808080"}

HDR_FILL = PatternFill("solid", fgColor="305496")
HDR_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")


def norm_status(s: str) -> str:
    return (s or "").strip().lower().replace("-", "_").replace(" ", "_")


def main() -> int:
    if not CSV_PATH.exists():
        sys.exit(f"missing {CSV_PATH}")
    with CSV_PATH.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    header, body = rows[0], rows[1:]

    wb = Workbook()
    ws = wb.active
    ws.title = "TODO"

    col = {name: i for i, name in enumerate(header)}
    status_i = col.get("status")
    prio_i = col.get("priority")

    # header row
    for c, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = TOP

    # body
    for r, record in enumerate(body, start=2):
        for c, val in enumerate(record):
            cell = ws.cell(row=r, column=c + 1, value=val)
            cell.border = BORDER
            cell.alignment = WRAP if header[c] in ("title", "comment") else TOP
            if c == status_i:
                fill, fg = STATUS_STYLE.get(norm_status(val), ("FFFFFF", "000000"))
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(bold=True, color=fg)
                cell.value = val.strip().lower()
            elif c == prio_i and val.strip().lower() in PRIORITY_FONT:
                cell.font = Font(color=PRIORITY_FONT[val.strip().lower()], bold=True)

    # widths
    widths = {"id": 6, "title": 46, "type": 9, "complexity": 11, "priority": 9,
              "date_created": 13, "status": 12, "date_closed": 12, "comment": 70}
    for c, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 14)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(body) + 1}"

    # summary counts in a side note (top-right, beyond table)
    counts: dict[str, int] = {}
    for record in body:
        counts[norm_status(record[status_i])] = counts.get(norm_status(record[status_i]), 0) + 1
    note_col = len(header) + 2
    ws.cell(row=1, column=note_col, value="Summary").font = Font(bold=True)
    for i, key in enumerate(("open", "in_progress", "done"), start=2):
        ws.cell(row=i, column=note_col, value=f"{key}: {counts.get(key, 0)}")

    wb.save(XLSX_PATH)
    total = len(body)
    print(f"wrote {XLSX_PATH}  ({total} rows: "
          f"open {counts.get('open', 0)} / in_progress {counts.get('in_progress', 0)} / "
          f"done {counts.get('done', 0)})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
